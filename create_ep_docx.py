from docx import Document

# 读取run_first_version_experiments.py的内容
with open('run_first_version_experiments.py', 'r', encoding='utf-8') as f:
    python_code = f.read()

# 创建文档
doc = Document()

# 添加标题
doc.add_heading('EP', level=1)

# 添加代码内容
doc.add_paragraph(python_code)

# 保存文档
doc.save('EP.docx')
print('Word document created: EP.docx')
